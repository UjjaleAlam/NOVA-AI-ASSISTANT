"""
Office Agent - Phase 12
Local document/spreadsheet/presentation manipulation.
Uses python-docx, openpyxl, python-pptx, and pandas.
No cloud APIs, no costs, fully open source.
"""

import os
import json
from typing import List, Dict, Any, Optional
from pathlib import Path
from dataclasses import dataclass
from enum import Enum

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

try:
    from pptx import Presentation
    from pptx.util import Inches as PptxInches, Pt as PptxPt
    from pptx.dml.color import RGBColor as PptxRGBColor
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False


class DocumentType(Enum):
    WORD = "word"
    EXCEL = "excel"
    POWERPOINT = "powerpoint"
    CSV = "csv"
    JSON = "json"
    MARKDOWN = "markdown"


@dataclass
class OfficeResult:
    success: bool
    message: str
    file_path: str = ""
    data: Any = None


class OfficeAgent:
    def __init__(self):
        self.supported_formats = {
            DocumentType.WORD: [".docx"],
            DocumentType.EXCEL: [".xlsx", ".xls"],
            DocumentType.POWERPOINT: [".pptx"],
            DocumentType.CSV: [".csv"],
            DocumentType.JSON: [".json"],
            DocumentType.MARKDOWN: [".md", ".markdown"],
        }

    # ==========================================
    # WORD DOCUMENTS
    # ==========================================

    def create_word_doc(self, title: str = "", content: str = "", file_path: str = "document.docx") -> OfficeResult:
        if not DOCX_AVAILABLE:
            return OfficeResult(False, "python-docx not installed. Run: pip install python-docx")

        try:
            doc = Document()

            if title:
                heading = doc.add_heading(title, 0)
                heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

            if content:
                for para in content.split("\n\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())

            doc.save(file_path)
            return OfficeResult(True, f"Created Word document: {file_path}", file_path)

        except Exception as e:
            return OfficeResult(False, f"Error creating Word doc: {e}")

    def read_word_doc(self, file_path: str) -> OfficeResult:
        if not DOCX_AVAILABLE:
            return OfficeResult(False, "python-docx not installed")

        try:
            doc = Document(file_path)
            content = []
            for para in doc.paragraphs:
                if para.text.strip():
                    content.append(para.text)
            return OfficeResult(True, "Document read successfully", file_path, "\n\n".join(content))

        except Exception as e:
            return OfficeResult(False, f"Error reading Word doc: {e}")

    def edit_word_doc(self, file_path: str, operations: List[Dict]) -> OfficeResult:
        if not DOCX_AVAILABLE:
            return OfficeResult(False, "python-docx not installed")

        try:
            doc = Document(file_path)

            for op in operations:
                op_type = op.get("type")

                if op_type == "add_heading":
                    doc.add_heading(op["text"], level=op.get("level", 1))

                elif op_type == "add_paragraph":
                    p = doc.add_paragraph(op["text"])
                    if op.get("bold"):
                        for run in p.runs:
                            run.bold = True
                    if op.get("alignment"):
                        p.alignment = getattr(WD_ALIGN_PARAGRAPH, op["alignment"].upper())

                elif op_type == "add_table":
                    rows = op.get("rows", 2)
                    cols = op.get("cols", 2)
                    data = op.get("data", [])
                    table = doc.add_table(rows=rows, cols=cols)
                    for i, row_data in enumerate(data):
                        for j, cell_text in enumerate(row_data):
                            if i < rows and j < cols:
                                table.cell(i, j).text = str(cell_text)

                elif op_type == "add_page_break":
                    doc.add_page_break()

            doc.save(file_path)
            return OfficeResult(True, f"Updated Word document: {file_path}", file_path)

        except Exception as e:
            return OfficeResult(False, f"Error editing Word doc: {e}")

    def convert_word_to_text(self, file_path: str) -> OfficeResult:
        return self.read_word_doc(file_path)

    # ==========================================
    # EXCEL SPREADSHEETS
    # ==========================================

    def create_excel(self, file_path: str = "spreadsheet.xlsx", sheets: Dict = None) -> OfficeResult:
        if not OPENPYXL_AVAILABLE:
            return OfficeResult(False, "openpyxl not installed. Run: pip install openpyxl")

        try:
            wb = openpyxl.Workbook()

            if sheets:
                first = True
                for sheet_name, sheet_data in sheets.items():
                    if first:
                        ws = wb.active
                        ws.title = sheet_name
                        first = False
                    else:
                        ws = wb.create_sheet(sheet_name)

                    # Write headers
                    headers = sheet_data.get("headers", [])
                    for col, header in enumerate(headers, 1):
                        cell = ws.cell(row=1, column=col, value=header)
                        cell.font = Font(bold=True)
                        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                        cell.font = Font(bold=True, color="FFFFFF")

                    # Write data
                    for row_idx, row_data in enumerate(sheet_data.get("data", []), 2):
                        for col_idx, value in enumerate(row_data, 1):
                            ws.cell(row=row_idx, column=col_idx, value=value)

                    # Auto-fit columns
                    for col in ws.columns:
                        max_len = max(len(str(cell.value or "")) for cell in col)
                        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 50)
            else:
                ws = wb.active
                ws.title = "Sheet1"

            wb.save(file_path)
            return OfficeResult(True, f"Created Excel file: {file_path}", file_path)

        except Exception as e:
            return OfficeResult(False, f"Error creating Excel: {e}")

    def read_excel(self, file_path: str, sheet_name: str = None) -> OfficeResult:
        if not OPENPYXL_AVAILABLE:
            return OfficeResult(False, "openpyxl not installed")

        try:
            wb = openpyxl.load_workbook(file_path)
            ws = wb[sheet_name] if sheet_name else wb.active

            data = []
            headers = []

            for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
                if row_idx == 1:
                    headers = [str(c) if c else f"Col{i}" for i, c in enumerate(row, 1)]
                else:
                    row_dict = {headers[i] if i < len(headers) else f"Col{i+1}": v for i, v in enumerate(row)}
                    data.append(row_dict)

            return OfficeResult(True, f"Read Excel file: {file_path}", file_path, {
                "sheet": ws.title,
                "headers": headers,
                "rows": len(data),
                "data": data
            })

        except Exception as e:
            return OfficeResult(False, f"Error reading Excel: {e}")

    def edit_excel(self, file_path: str, operations: List[Dict]) -> OfficeResult:
        if not OPENPYXL_AVAILABLE:
            return OfficeResult(False, "openpyxl not installed")

        try:
            wb = openpyxl.load_workbook(file_path)

            for op in operations:
                op_type = op.get("type")
                sheet_name = op.get("sheet", wb.active.title)
                ws = wb[sheet_name]

                if op_type == "write_cell":
                    ws.cell(row=op["row"], column=op["col"], value=op["value"])

                elif op_type == "write_row":
                    for col_idx, value in enumerate(op["values"], op.get("start_col", 1)):
                        ws.cell(row=op["row"], column=col_idx, value=value)

                elif op_type == "append_row":
                    ws.append(op["values"])

                elif op_type == "delete_row":
                    ws.delete_rows(op["row"], op.get("count", 1))

                elif op_type == "format_cell":
                    cell = ws.cell(row=op["row"], column=op["col"])
                    if "bold" in op:
                        cell.font = Font(bold=op["bold"])
                    if "color" in op:
                        cell.font = Font(color=op["color"])
                    if "fill" in op:
                        cell.fill = PatternFill(start_color=op["fill"], end_color=op["fill"], fill_type="solid")
                    if "alignment" in op:
                        cell.alignment = Alignment(horizontal=op["alignment"])

                elif op_type == "set_column_width":
                    ws.column_dimensions[get_column_letter(op["col"])].width = op["width"]

            wb.save(file_path)
            return OfficeResult(True, f"Updated Excel file: {file_path}", file_path)

        except Exception as e:
            return OfficeResult(False, f"Error editing Excel: {e}")

    def excel_to_csv(self, file_path: str, output_path: str = None, sheet_name: str = None) -> OfficeResult:
        if not PANDAS_AVAILABLE:
            return OfficeResult(False, "pandas not installed. Run: pip install pandas")

        try:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            output_path = output_path or file_path.replace(".xlsx", ".csv")
            df.to_csv(output_path, index=False)
            return OfficeResult(True, f"Converted to CSV: {output_path}", output_path)

        except Exception as e:
            return OfficeResult(False, f"Error converting to CSV: {e}")

    def csv_to_excel(self, file_path: str, output_path: str = None) -> OfficeResult:
        if not PANDAS_AVAILABLE:
            return OfficeResult(False, "pandas not installed")

        try:
            df = pd.read_csv(file_path)
            output_path = output_path or file_path.replace(".csv", ".xlsx")
            df.to_excel(output_path, index=False)
            return OfficeResult(True, f"Converted to Excel: {output_path}", output_path)

        except Exception as e:
            return OfficeResult(False, f"Error converting to Excel: {e}")

    # ==========================================
    # POWERPOINT
    # ==========================================

    def create_powerpoint(self, file_path: str = "presentation.pptx", slides: List[Dict] = None) -> OfficeResult:
        if not PPTX_AVAILABLE:
            return OfficeResult(False, "python-pptx not installed. Run: pip install python-pptx")

        try:
            prs = Presentation()

            if slides:
                for slide_data in slides:
                    layout = prs.slide_layouts[slide_data.get("layout", 1)]
                    slide = prs.slides.add_slide(layout)

                    # Title
                    if "title" in slide_data:
                        title_shape = slide.shapes.title
                        if title_shape:
                            title_shape.text = slide_data["title"]

                    # Content
                    if "content" in slide_data:
                        for shape in slide.shapes:
                            if shape.has_text_frame and shape != slide.shapes.title:
                                shape.text = slide_data["content"]
                                break

                    # Bullets
                    if "bullets" in slide_data:
                        for shape in slide.shapes:
                            if shape.has_text_frame and shape != slide.shapes.title:
                                tf = shape.text_frame
                                tf.clear()
                                for i, bullet in enumerate(slide_data["bullets"]):
                                    if i == 0:
                                        tf.text = bullet
                                    else:
                                        p = tf.add_paragraph()
                                        p.text = bullet
                                        p.level = 0
                                break
            else:
                # Default blank presentation
                slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

            prs.save(file_path)
            return OfficeResult(True, f"Created PowerPoint: {file_path}", file_path)

        except Exception as e:
            return OfficeResult(False, f"Error creating PowerPoint: {e}")

    def read_powerpoint(self, file_path: str) -> OfficeResult:
        if not PPTX_AVAILABLE:
            return OfficeResult(False, "python-pptx not installed")

        try:
            prs = Presentation(file_path)
            slides_data = []

            for i, slide in enumerate(prs.slides):
                slide_data = {"slide": i + 1, "title": "", "content": []}
                for shape in slide.shapes:
                    if shape.has_text_frame:
                        text = shape.text_frame.text.strip()
                        if text:
                            if not slide_data["title"]:
                                slide_data["title"] = text
                            else:
                                slide_data["content"].append(text)
                slides_data.append(slide_data)

            return OfficeResult(True, f"Read PowerPoint: {file_path}", file_path, {"slides": slides_data})

        except Exception as e:
            return OfficeResult(False, f"Error reading PowerPoint: {e}")

    # ==========================================
    # CSV / JSON / MARKDOWN
    # ==========================================

    def read_csv(self, file_path: str) -> OfficeResult:
        if not PANDAS_AVAILABLE:
            return OfficeResult(False, "pandas not installed")

        try:
            df = pd.read_csv(file_path)
            return OfficeResult(True, f"Read CSV: {file_path}", file_path, {
                "columns": list(df.columns),
                "rows": len(df),
                "data": df.to_dict("records")
            })
        except Exception as e:
            return OfficeResult(False, f"Error reading CSV: {e}")

    def write_csv(self, data: List[Dict], file_path: str = "data.csv") -> OfficeResult:
        if not PANDAS_AVAILABLE:
            return OfficeResult(False, "pandas not installed")

        try:
            df = pd.DataFrame(data)
            df.to_csv(file_path, index=False)
            return OfficeResult(True, f"Wrote CSV: {file_path}", file_path)
        except Exception as e:
            return OfficeResult(False, f"Error writing CSV: {e}")

    def read_json(self, file_path: str) -> OfficeResult:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                data = json.load(f)
            return OfficeResult(True, f"Read JSON: {file_path}", file_path, data)
        except Exception as e:
            return OfficeResult(False, f"Error reading JSON: {e}")

    def write_json(self, data: Any, file_path: str = "data.json") -> OfficeResult:
        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                json.dump(data, f, indent=2)
            return OfficeResult(True, f"Wrote JSON: {file_path}", file_path)
        except Exception as e:
            return OfficeResult(False, f"Error writing JSON: {e}")

    def read_markdown(self, file_path: str) -> OfficeResult:
        try:
            with open(file_path, 'r', encoding='utf-8') as f:
                content = f.read()
            return OfficeResult(True, f"Read Markdown: {file_path}", file_path, content)
        except Exception as e:
            return OfficeResult(False, f"Error reading Markdown: {e}")

    def write_markdown(self, content: str, file_path: str = "document.md") -> OfficeResult:
        try:
            with open(file_path, 'w', encoding='utf-8') as f:
                f.write(content)
            return OfficeResult(True, f"Wrote Markdown: {file_path}", file_path)
        except Exception as e:
            return OfficeResult(False, f"Error writing Markdown: {e}")

    # ==========================================
    # UNIVERSAL INTERFACE
    # ==========================================

    def create_document(self, doc_type: DocumentType, file_path: str, **kwargs) -> OfficeResult:
        if doc_type == DocumentType.WORD:
            return self.create_word_doc(kwargs.get("title", ""), kwargs.get("content", ""), file_path)
        elif doc_type == DocumentType.EXCEL:
            return self.create_excel(file_path, kwargs.get("sheets"))
        elif doc_type == DocumentType.POWERPOINT:
            return self.create_powerpoint(file_path, kwargs.get("slides"))
        elif doc_type == DocumentType.CSV:
            return self.write_csv(kwargs.get("data", []), file_path)
        elif doc_type == DocumentType.JSON:
            return self.write_json(kwargs.get("data", {}), file_path)
        elif doc_type == DocumentType.MARKDOWN:
            return self.write_markdown(kwargs.get("content", ""), file_path)
        return OfficeResult(False, f"Unsupported document type: {doc_type}")

    def read_document(self, file_path: str) -> OfficeResult:
        ext = Path(file_path).suffix.lower()
        if ext in [".docx"]:
            return self.read_word_doc(file_path)
        elif ext in [".xlsx", ".xls"]:
            return self.read_excel(file_path)
        elif ext in [".pptx"]:
            return self.read_powerpoint(file_path)
        elif ext in [".csv"]:
            return self.read_csv(file_path)
        elif ext in [".json"]:
            return self.read_json(file_path)
        elif ext in [".md", ".markdown"]:
            return self.read_markdown(file_path)
        return OfficeResult(False, f"Unsupported file type: {ext}")

    def get_document_info(self, file_path: str) -> OfficeResult:
        """Get metadata about a document."""
        if not os.path.exists(file_path):
            return OfficeResult(False, "File not found")

        stat = os.stat(file_path)
        ext = Path(file_path).suffix.lower()

        info = {
            "path": file_path,
            "name": Path(file_path).name,
            "extension": ext,
            "size": stat.st_size,
            "modified": stat.st_mtime,
            "type": "unknown"
        }

        for doc_type, extensions in self.supported_formats.items():
            if ext in extensions:
                info["type"] = doc_type.value
                break

        return OfficeResult(True, "File info retrieved", file_path, info)


# Global instance
office_agent = OfficeAgent()


# Convenience functions
def create_word(title: str = "", content: str = "", path: str = "document.docx"):
    return office_agent.create_word_doc(title, content, path)

def create_excel(path: str = "spreadsheet.xlsx", sheets: Dict = None):
    return office_agent.create_excel(path, sheets)

def create_ppt(path: str = "presentation.pptx", slides: List = None):
    return office_agent.create_powerpoint(path, slides)

def read_file(path: str):
    return office_agent.read_document(path)

def get_file_info(path: str):
    return office_agent.get_document_info(path)


if __name__ == "__main__":
    agent = OfficeAgent()

    # Test Word
    if DOCX_AVAILABLE:
        result = agent.create_word_doc("Test Document", "This is a test paragraph.\n\nAnother paragraph.")
        print("Word:", result.message)

        result = agent.read_word_doc("document.docx")
        print("Read Word:", result.data[:100] if result.data else "None")

    # Test Excel
    if OPENPYXL_AVAILABLE:
        sheets = {
            "Users": {
                "headers": ["Name", "Email", "Role"],
                "data": [
                    ["Alice", "alice@example.com", "Admin"],
                    ["Bob", "bob@example.com", "User"],
                    ["Carol", "carol@example.com", "Editor"]
                ]
            },
            "Stats": {
                "headers": ["Metric", "Value"],
                "data": [
                    ["Total Users", 3],
                    ["Active", 2],
                    ["Inactive", 1]
                ]
            }
        }
        result = agent.create_excel("test.xlsx", sheets)
        print("Excel:", result.message)

        result = agent.read_excel("test.xlsx")
        print("Read Excel:", result.data)

    # Test CSV
    if PANDAS_AVAILABLE:
        data = [{"name": "Alice", "age": 30}, {"name": "Bob", "age": 25}]
        result = agent.write_csv(data, "test.csv")
        print("CSV:", result.message)

        result = agent.read_csv("test.csv")
        print("Read CSV:", result.data)

    # Test JSON
    result = agent.write_json({"project": "Nova", "version": "1.0"}, "test.json")
    print("JSON:", result.message)

    result = agent.read_json("test.json")
    print("Read JSON:", result.data)

    # Test Markdown
    md = "# Test\n\nThis is **markdown** content."
    result = agent.write_markdown(md, "test.md")
    print("Markdown:", result.message)

    result = agent.read_markdown("test.md")
    print("Read MD:", result.data[:50])